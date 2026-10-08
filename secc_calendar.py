"""
SECC trade fair calendar updater - JTM Asia
-------------------------------------------
1. Reads the event list at https://secc.com.vn/events
2. Merges it into data/events.json (the master list - past events are kept)
3. Writes:
     docs/secc_trade_fairs.ics   -> the calendar Outlook subscribes to
     docs/secc_trade_fairs.xlsx  -> same events as a spreadsheet
     docs/events.json            -> read by the Power Automate flow

Run it on your own computer:
    python secc_calendar.py
Test with a saved copy of the page (no internet needed):
    python secc_calendar.py --html saved_page.html

Safety:
- If the website returns 0 events (site down, layout changed), nothing is
  overwritten and the script stops with an error.
- The master list is saved through a temp file + backup, so a crash
  half-way never corrupts it.
- Each event keeps the same ID forever, so Outlook updates an event when
  dates change instead of creating a duplicate.
"""

import argparse
import hashlib
import json
import os
import re
import shutil
import sys
import time
from datetime import date, datetime, timedelta, timezone

import requests
from bs4 import BeautifulSoup

URL = "https://secc.com.vn/events"
VENUE = "SECC, 799 Nguyen Van Linh, Tan My Ward, Ho Chi Minh City, Vietnam"
CAL_NAME = "SECC Trade Fairs (HCMC)"

HERE = os.path.dirname(os.path.abspath(__file__))
DATA_FILE = os.path.join(HERE, "data", "events.json")
ICS_FILE = os.path.join(HERE, "docs", "secc_trade_fairs.ics")
XLSX_FILE = os.path.join(HERE, "docs", "secc_trade_fairs.xlsx")

HEADERS = {
    "User-Agent": ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                   "(KHTML, like Gecko) Chrome/126.0 Safari/537.36"),
    "Accept-Language": "en-US,en;q=0.9,vi;q=0.8",
}

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


# ---------------------------------------------------------------- fetch
def fetch_html(saved_file=None):
    if saved_file:
        with open(saved_file, encoding="utf-8") as f:
            return f.read()
    last_err = None
    for attempt in range(1, 4):
        try:
            r = requests.get(URL, headers=HEADERS, timeout=60)
            r.raise_for_status()
            return r.text
        except Exception as e:  # retry on network trouble
            last_err = e
            print(f"  attempt {attempt} failed: {e}")
            time.sleep(10 * attempt)
    raise SystemExit(f"Could not load {URL}: {last_err}")


# ---------------------------------------------------------------- dates
def parse_dates(text):
    """Turn SECC date text into (start, end) dates.
    Handles: '14 - 17.10.2026', '28.10 - 02.11.2026',
             '30.12.2026 - 02.01.2027', '14.10.2026'"""
    t = re.sub(r"\s+", "", text.replace("\u2013", "-").replace("\u2014", "-"))
    t = t.replace("/", ".")
    parts = t.split("-")
    if len(parts) == 1:
        d, m, y = map(int, parts[0].split("."))
        one = date(y, m, d)
        return one, one
    left, right = parts[0], parts[-1]
    rd, rm, ry = map(int, right.split("."))
    lp = [int(x) for x in left.split(".") if x]
    ld = lp[0]
    lm = lp[1] if len(lp) > 1 else rm
    ly = lp[2] if len(lp) > 2 else (ry - 1 if lm > rm else ry)
    return date(ly, lm, ld), date(ry, rm, rd)


def nice_range(s, e):
    if s == e:
        return f"{s.day} {MONTHS[s.month-1]} {s.year}"
    if s.year == e.year and s.month == e.month:
        return f"{s.day}-{e.day} {MONTHS[s.month-1]} {s.year}"
    if s.year == e.year:
        return f"{s.day} {MONTHS[s.month-1]} - {e.day} {MONTHS[e.month-1]} {s.year}"
    return f"{s.day} {MONTHS[s.month-1]} {s.year} - {e.day} {MONTHS[e.month-1]} {e.year}"


# ---------------------------------------------------------------- parse
def parse_events(html):
    soup = BeautifulSoup(html, "html.parser")
    events = []
    for block in soup.select(".table-events td.has-bg .block"):
        try:
            a = block.find("a", href=True)
            h3 = block.find("h3")
            if not h3:
                continue
            name = h3.get_text(" ", strip=True)
            link = a["href"].strip() if a else URL

            # subtitle lines = <p> tags directly in the block (not the icon rows)
            subtitles = [p.get_text(" ", strip=True) for p in block.find_all("p", recursive=False)]
            subtitles = [s for s in subtitles if s]

            def icon_texts(icon):
                out = []
                for i in block.select(f"i.{icon}"):
                    span = i.find_next_sibling("span")
                    if span and span.get_text(strip=True):
                        out.append(span.get_text(", ", strip=True))
                return out

            date_txt = (icon_texts("fa-calendar") or [""])[0]
            if not date_txt:
                print(f"  skipped (no date): {name}")
                continue
            start, end = parse_dates(date_txt)
            hours = (icon_texts("fa-clock-o") or [""])[0].replace(" ", "")
            halls = icon_texts("fa-map-marker")

            extra = {}
            for p in block.select(".col-md-6 p"):
                la = p.find("a", href=True)
                if la:
                    extra[la.get_text(" ", strip=True)] = la["href"].strip()

            m = re.search(r"-(\d{6,})/?$", link)
            uid_key = m.group(1) if m else f"{name}-{start.isoformat()}"

            events.append({
                "id": uid_key,
                "name": name,
                "subtitle": " / ".join(subtitles),
                "start": start.isoformat(),
                "end": end.isoformat(),
                "date_text": date_txt,
                "hours": hours,
                "halls": halls,
                "link": link,
                "brochure": extra.get("Brochure", ""),
                "report": extra.get("Post show report", ""),
            })
        except Exception as e:  # one bad row never stops the run
            print(f"  could not read one event: {e}")
    return events


# ---------------------------------------------------------------- storage
def load_store():
    if os.path.exists(DATA_FILE):
        with open(DATA_FILE, encoding="utf-8") as f:
            return json.load(f)
    return {"events": {}}


def save_store(store):
    os.makedirs(os.path.dirname(DATA_FILE), exist_ok=True)
    if os.path.exists(DATA_FILE):
        shutil.copy2(DATA_FILE, DATA_FILE + ".bak")
    tmp = DATA_FILE + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(store, f, ensure_ascii=False, indent=2)
    os.replace(tmp, DATA_FILE)


CONTENT_FIELDS = ["name", "subtitle", "start", "end", "hours", "halls", "link", "brochure", "report"]


def fingerprint(ev):
    return hashlib.sha1(json.dumps([ev.get(k) for k in CONTENT_FIELDS],
                                   ensure_ascii=False).encode()).hexdigest()


def merge(store, scraped, today):
    stats = {"new": 0, "changed": 0, "same": 0, "missing": 0}
    seen = set()
    for ev in scraped:
        seen.add(ev["id"])
        old = store["events"].get(ev["id"])
        fp = fingerprint(ev)
        if old is None:
            ev.update(first_seen=today, sequence=0, fp=fp)
            stats["new"] += 1
        else:
            ev["first_seen"] = old.get("first_seen", today)
            ev["sequence"] = old.get("sequence", 0)
            if old.get("fp") != fp or old.get("not_listed_since"):
                ev["sequence"] += 1
                stats["changed"] += 1
            else:
                stats["same"] += 1
            ev["fp"] = fp
        ev["last_seen"] = today
        ev.pop("not_listed_since", None)
        store["events"][ev["id"]] = ev

    # upcoming events that vanished from the site: keep, but flag them
    for key, ev in store["events"].items():
        if key in seen or ev["end"] < today:
            continue
        if not ev.get("not_listed_since"):
            ev["not_listed_since"] = today
            ev["sequence"] = ev.get("sequence", 0) + 1
        stats["missing"] += 1
    store["last_checked"] = today
    return stats


# ---------------------------------------------------------------- ics
def ics_escape(s):
    return (s.replace("\\", "\\\\").replace(";", "\\;")
             .replace(",", "\\,").replace("\n", "\\n"))


def fold(line):
    """Fold lines at 75 bytes (iCalendar rule) without breaking UTF-8 letters."""
    out, cur = [], ""
    for ch in line:
        limit = 75 if not out else 74
        if len((cur + ch).encode("utf-8")) > limit:
            out.append(cur)
            cur = ch
        else:
            cur += ch
    out.append(cur)
    return "\r\n ".join(out)


def description(ev, checked):
    s, e = date.fromisoformat(ev["start"]), date.fromisoformat(ev["end"])
    lines = []
    if ev.get("not_listed_since"):
        lines.append(f"NOTE: no longer listed on the SECC website since {ev['not_listed_since']} - "
                     "please check if it was moved or cancelled.")
        lines.append("")
    if ev.get("subtitle"):
        lines.append(ev["subtitle"])
        lines.append("")
    lines.append(f"Dates: {nice_range(s, e)}")
    if ev.get("hours"):
        lines.append(f"Hours: {ev['hours']}")
    if ev.get("halls"):
        lines.append("Halls: " + ", ".join(ev["halls"]))
    lines.append(f"Venue: {VENUE}")
    lines.append("")
    lines.append(f"Event page: {ev['link']}")
    if ev.get("brochure"):
        lines.append(f"Brochure: {ev['brochure']}")
    if ev.get("report"):
        lines.append(f"Post show report: {ev['report']}")
    lines.append("")
    lines.append(f"Source: SECC event calendar, last checked {checked}")
    return "\n".join(lines)


def location(ev):
    halls = ", ".join(ev.get("halls") or [])
    return f"{VENUE} | {halls}" if halls else VENUE


def write_ics(store):
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    checked = store.get("last_checked", "")
    rows = [
        "BEGIN:VCALENDAR", "VERSION:2.0", "CALSCALE:GREGORIAN", "METHOD:PUBLISH",
        "PRODID:-//JTM Asia//SECC Trade Fairs//EN",
        f"X-WR-CALNAME:{CAL_NAME}",
        "X-WR-CALDESC:Trade fairs at SECC Ho Chi Minh City - auto-updated from secc.com.vn",
        "X-WR-TIMEZONE:Asia/Ho_Chi_Minh",
        "REFRESH-INTERVAL;VALUE=DURATION:P1D",
        "X-PUBLISHED-TTL:P1D",
    ]
    for ev in sorted(store["events"].values(), key=lambda x: (x["start"], x["name"])):
        s = date.fromisoformat(ev["start"])
        e = date.fromisoformat(ev["end"]) + timedelta(days=1)  # all-day end is exclusive
        rows += [
            "BEGIN:VEVENT",
            f"UID:secc-{ev['id']}@jtmasia.com",
            f"DTSTAMP:{stamp}",
            f"DTSTART;VALUE=DATE:{s:%Y%m%d}",
            f"DTEND;VALUE=DATE:{e:%Y%m%d}",
            f"SUMMARY:{ics_escape(ev['name'])}",
            f"LOCATION:{ics_escape(location(ev))}",
            f"DESCRIPTION:{ics_escape(description(ev, checked))}",
            f"URL:{ev['link']}",
            f"SEQUENCE:{ev.get('sequence', 0)}",
            "STATUS:" + ("TENTATIVE" if ev.get("not_listed_since") else "CONFIRMED"),
            "TRANSP:TRANSPARENT",
            "END:VEVENT",
        ]
    rows.append("END:VCALENDAR")
    os.makedirs(os.path.dirname(ICS_FILE), exist_ok=True)
    tmp = ICS_FILE + ".tmp"
    with open(tmp, "w", encoding="utf-8", newline="") as f:
        f.write("\r\n".join(fold(r) for r in rows) + "\r\n")
    os.replace(tmp, ICS_FILE)


# ---------------------------------------------------------------- json for Power Automate
def html_escape(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def write_json(store):
    """docs/events.json - read by the Power Automate flow.
    'version' only goes up when an event really changes, so the flow
    skips events that are the same as last time."""
    checked = store.get("last_checked", "")
    out = []
    for ev in sorted(store["events"].values(), key=lambda x: (x["start"], x["name"])):
        s = date.fromisoformat(ev["start"])
        e = date.fromisoformat(ev["end"]) + timedelta(days=1)  # all-day end is exclusive
        desc = description(ev, checked)
        body_html = "<br>".join(
            (f'<a href="{l.split(": ", 1)[1]}">{html_escape(l)}</a>'
             if re.match(r"^(Event page|Brochure|Post show report): https?://", l) else html_escape(l))
            for l in desc.split("\n"))
        out.append({
            "key": f"SECC-{ev['id']}",
            "name": ev["name"],
            "start": ev["start"],
            "end": ev["end"],
            "start_time": f"{s.isoformat()}T00:00:00",
            "end_time": f"{e.isoformat()}T00:00:00",
            "hours": ev.get("hours", ""),
            "halls": ", ".join(ev.get("halls") or []),
            "location": location(ev),
            "link": ev["link"],
            "status": ("Not listed since " + ev["not_listed_since"]) if ev.get("not_listed_since")
                      else ("Past" if ev["end"] < checked else "Upcoming"),
            "version": ev.get("sequence", 0),
            "body_html": body_html,
        })
    path = os.path.join(HERE, "docs", "events.json")
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump({"last_checked": checked, "events": out}, f, ensure_ascii=False, indent=1)
    os.replace(tmp, path)


# ---------------------------------------------------------------- xlsx
def write_xlsx(store):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print("  openpyxl not installed - skipping spreadsheet")
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "SECC events"
    head = ["Subject", "Start Date", "End Date", "Hours", "Halls", "Description",
            "Event page", "Brochure", "Status", "First seen", "Last seen"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for ev in sorted(store["events"].values(), key=lambda x: (x["start"], x["name"])):
        status = (f"Not listed since {ev['not_listed_since']}" if ev.get("not_listed_since")
                  else ("Past" if ev["end"] < store["last_checked"] else "Upcoming"))
        ws.append([ev["name"], date.fromisoformat(ev["start"]), date.fromisoformat(ev["end"]),
                   ev.get("hours", ""), ", ".join(ev.get("halls") or []), ev.get("subtitle", ""),
                   ev["link"], ev.get("brochure", ""), status,
                   ev.get("first_seen", ""), ev.get("last_seen", "")])
    widths = [42, 12, 12, 12, 30, 60, 50, 50, 24, 12, 12]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = row[2].number_format = "yyyy-mm-dd"
        row[5].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    tmp = XLSX_FILE.replace(".xlsx", ".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, XLSX_FILE)


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--html", help="use a saved copy of the SECC events page instead of the website")
    ap.add_argument("--today", help="pretend today is YYYY-MM-DD (for testing)")
    args = ap.parse_args()
    today = args.today or date.today().isoformat()

    print(f"Reading {'saved page ' + args.html if args.html else URL} ...")
    scraped = parse_events(fetch_html(args.html))
    print(f"  found {len(scraped)} events on the page")
    if not scraped:
        raise SystemExit("ERROR: 0 events found - the site may be down or its layout changed. "
                         "Nothing was overwritten.")

    store = load_store()
    stats = merge(store, scraped, today)
    save_store(store)
    write_ics(store)
    write_json(store)
    write_xlsx(store)

    print(f"  new: {stats['new']}  changed: {stats['changed']}  unchanged: {stats['same']}  "
          f"upcoming but no longer listed: {stats['missing']}")
    print(f"  total events in calendar: {len(store['events'])}")
    print(f"Saved: {ICS_FILE}")
    print(f"Saved: {XLSX_FILE}")


if __name__ == "__main__":
    main()
