#!/usr/bin/env python3
"""
WTC Binh Duong event calendar.

Reads https://wtcbinhduong.vn/events/ (and the English page of each event),
merges it into data/wtc_events.json and writes:
  docs/wtc_binh_duong.ics    - calendar for Outlook "Subscribe from web"
  docs/wtc_events.json       - upcoming events for Power Automate
  docs/WTC Binh Duong Trade Fairs.xlsx - the full list as a spreadsheet

Safe to stop and run again: progress is saved after every event.
Nothing is overwritten if the site returns no events.
"""
import json
import re
import sys
import time
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import requests
from bs4 import BeautifulSoup

BASE = "https://wtcbinhduong.vn"
LIST_URL = BASE + "/events/"
LIST_PAGE_URL = BASE + "/events/page/{n}/"
MAX_PAGES = 40
DELAY = 1.0  # seconds between requests, to be polite

HERE = Path(__file__).resolve().parent
DATA_FILE = HERE / "data" / "wtc_events.json"
DOCS = HERE / "docs"
ICS_FILE = DOCS / "wtc_binh_duong.ics"
PA_FILE = DOCS / "wtc_events.json"
XLSX_FILE = DOCS / "WTC Binh Duong Trade Fairs.xlsx"

VN_TZ = timezone(timedelta(hours=7))
TODAY = datetime.now(VN_TZ).date()
# First run: keep events from 1 January of last year onwards.
HISTORY_FROM = date(TODAY.year - 1, 1, 1)

HEADERS = {
    "User-Agent": "Mozilla/5.0 (compatible; JTM-Asia-calendar/1.0; +mailto:julia.hoang@jtmasia.com)",
    "Accept-Language": "vi,en;q=0.8",
}

# Venue name (any language) -> (English venue name, English address)
VENUES = [
    ("bcec", "Binh Duong Convention & Exhibition Center (BCEC)",
     "B11 Hung Vuong Street, Binh Duong Ward, Ho Chi Minh City, Vietnam"),
    ("hội nghị", "Binh Duong Convention & Exhibition Center (BCEC)",
     "B11 Hung Vuong Street, Binh Duong Ward, Ho Chi Minh City, Vietnam"),
    ("convention", "Binh Duong Convention & Exhibition Center (BCEC)",
     "B11 Hung Vuong Street, Binh Duong Ward, Ho Chi Minh City, Vietnam"),
    ("wtc expo", "WTC EXPO International Exhibition Center",
     "A19 Hung Vuong Street, Binh Duong Ward, Ho Chi Minh City, Vietnam"),
]
DEFAULT_ADDRESS = "WTC Binh Duong New City, Hung Vuong Street, Binh Duong Ward, Ho Chi Minh City, Vietnam"

DATE_RE = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4})")


# ---------------------------------------------------------------- helpers
session = requests.Session()
session.headers.update(HEADERS)


def get(url):
    """Download a page, with 3 tries."""
    last = None
    for attempt in range(3):
        try:
            r = session.get(url, timeout=30)
            if r.status_code == 404:
                return None
            r.raise_for_status()
            time.sleep(DELAY)
            return r.text
        except requests.RequestException as e:
            last = e
            print(f"  retry {attempt + 1}/3 for {url}: {e}")
            time.sleep(5 * (attempt + 1))
    raise RuntimeError(f"Could not download {url}: {last}")


def clean(text):
    return re.sub(r"\s+", " ", text or "").strip()


def to_date(d, m, y):
    try:
        return date(int(y), int(m), int(d))
    except ValueError:
        return None


def parse_dates(text):
    """'22/10/2026 - 24/10/2026' -> (start, end). One date -> (d, d)."""
    text = text or ""
    # '22 – 24/10/2026' (same month) or '30/09 – 02/10/2026' (two months)
    m = re.search(r"(?<![\d/])(\d{1,2})(?:/(\d{1,2}))?\s*[-–]\s*(\d{1,2})/(\d{1,2})/(\d{4})", text)
    if m:
        y = m.group(5)
        s = to_date(m.group(1), m.group(2) or m.group(4), y)
        e = to_date(m.group(3), m.group(4), y)
        if s and e and s <= e:
            return s, e
    found = [to_date(*m) for m in DATE_RE.findall(text)]
    found = [d for d in found if d]
    if not found:
        return None, None
    return found[0], found[-1] if found[-1] >= found[0] else found[0]


def venue_info(venue_text):
    low = (venue_text or "").lower()
    hall = ""
    m = re.match(r"\s*(hall\s+[a-z0-9]+)", low, re.I)
    if m:
        hall = m.group(1).title()
    for key, name, addr in VENUES:
        if key in low:
            return name, addr, hall
    return clean(venue_text), DEFAULT_ADDRESS, hall


def slug_of(url):
    return url.rstrip("/").split("/")[-1]


# ---------------------------------------------------------------- listing
def parse_listing(html):
    soup = BeautifulSoup(html, "html.parser")
    items = []
    for li in soup.select("li.event-item"):
        a = li.select_one(".title-28 a") or li.select_one('a[href*="/su-kien/"]')
        if not a or not a.get("href"):
            continue
        venue, when = "", ""
        for box in li.select(".info-box"):
            img = box.find("img")
            p = box.find("p")
            src = (img.get("src") if img else "") or ""
            if "location" in src:
                venue = clean(p.get_text() if p else "")
            elif "clock" in src:
                when = clean(p.get_text() if p else "")
        d, _ = parse_dates(when)
        hours = clean(DATE_RE.sub("", when)).strip(" -–")
        items.append({
            "url": a["href"].split("?")[0],
            "name_vn": clean(a.get_text()),
            "venue_vn": venue,
            "hours": hours,
            "list_date": d,
        })
    last_page = 1
    for a in soup.select("ul.page-numbers a.page-numbers"):
        if clean(a.get_text()).isdigit():
            last_page = max(last_page, int(clean(a.get_text())))
    return items, last_page


def scan_listing(cutoff):
    """All listing items starting on/after cutoff (pages are newest first)."""
    all_items, seen = [], set()
    n, last_page = 1, 1
    while n <= min(last_page, MAX_PAGES):
        url = LIST_URL if n == 1 else LIST_PAGE_URL.format(n=n)
        html = get(url)
        if html is None:
            break
        items, lp = parse_listing(html)
        last_page = max(last_page, lp)
        print(f"Page {n}: {len(items)} events")
        if not items:
            break
        for it in items:
            if it["url"] not in seen:
                seen.add(it["url"])
                all_items.append(it)
        dated = [it["list_date"] for it in items if it["list_date"]]
        if dated and max(dated) < cutoff:
            break
        n += 1
    return [it for it in all_items if not it["list_date"] or it["list_date"] >= cutoff]


# ---------------------------------------------------------------- detail
def parse_detail(html):
    soup = BeautifulSoup(html, "html.parser")
    out = {"name": "", "dates_text": "", "venue": "", "map": "", "en_url": "",
           "description": "", "website": "", "registration": "", "brochure": ""}
    h1 = soup.find("h1")
    out["name"] = clean(h1.get_text()) if h1 else ""

    for li in soup.select("li.list-item"):
        ps = li.select(".list-content > p")
        if len(ps) < 2:
            continue
        label = clean(ps[0].get_text()).lower()
        value = clean(ps[1].get_text())
        if label in ("ngày diễn ra", "event date") or "date" in label or "ngày" in label:
            out["dates_text"] = value
        elif label in ("địa điểm", "location"):
            out["venue"] = value
            m = li.select_one('a[href*="maps"]')
            out["map"] = m["href"] if m else ""

    lang = soup.select_one("a.language-child-link")
    if lang and "/en/" in (lang.get("href") or ""):
        out["en_url"] = lang["href"]

    content = soup.select_one(".mona-content")
    if content:
        # Description: the opening paragraphs, up to about 1500 characters.
        parts, size = [], 0
        for el in content.find_all(["p", "h2", "h3", "ul"], recursive=False):
            t = clean(el.get_text(" "))
            if not t:
                continue
            if el.name == "ul":
                t = "\n".join("- " + clean(x.get_text(" ")) for x in el.find_all("li"))
            parts.append(t)
            size += len(t)
            if size > 1500:
                break
        out["description"] = "\n\n".join(parts)
        if not out["dates_text"]:
            txt = content.get_text(" ")
            m = re.search(r"(Thời gian|Time|Date)\s*:?\s*([^\n]{0,60})", txt)
            if m:
                out["dates_text"] = m.group(2)
        # Links
        for a in content.find_all("a", href=True):
            href = a["href"].strip()
            low = href.lower()
            label = clean(a.get_text()).lower()
            if low.startswith(("mailto:", "tel:", "javascript")) or "maps" in low:
                continue
            if low.endswith(".pdf") or "brochure" in low or "brochure" in label:
                out["brochure"] = out["brochure"] or href
            elif "wtcbinhduong.vn" in low:
                continue
            elif re.search(r"regist|dang-?ky|đăng ký|register", low + " " + label):
                out["registration"] = out["registration"] or href
            elif not out["website"]:
                out["website"] = href
    return out


def fetch_event(item):
    vn = parse_detail(get(item["url"]) or "")
    en = parse_detail(get(vn["en_url"]) or "") if vn["en_url"] else None

    start, end = parse_dates(vn["dates_text"])
    if not start and en:
        start, end = parse_dates(en["dates_text"])
    if not start:
        start = end = item["list_date"]
    if not start:
        return None

    venue_name, address, hall = venue_info((en and en["venue"]) or vn["venue"] or item["venue_vn"])
    if not hall:
        hall = venue_info(item["venue_vn"])[2]

    return {
        "key": slug_of(item["url"]),
        "name": (en and en["name"]) or vn["name"] or item["name_vn"],
        "name_vn": vn["name"] or item["name_vn"],
        "start": start.isoformat(),
        "end": end.isoformat(),
        "hours": item["hours"],
        "venue": venue_name,
        "hall": hall,
        "address": address,
        "map": vn["map"] or (en and en["map"]) or "",
        "description": (en and en["description"]) or vn["description"],
        "description_vn": vn["description"] if en and en["description"] else "",
        "url": (en and vn["en_url"]) or item["url"],
        "url_vn": item["url"],
        "website": (en and en["website"]) or vn["website"],
        "registration": (en and en["registration"]) or vn["registration"],
        "brochure": (en and en["brochure"]) or vn["brochure"],
    }


# ---------------------------------------------------------------- merge
TRACKED = ("name", "start", "end", "hours", "venue", "hall", "status")


def load_data():
    if DATA_FILE.exists():
        return json.loads(DATA_FILE.read_text(encoding="utf-8"))
    return {"events": {}}


def save_data(data):
    DATA_FILE.parent.mkdir(parents=True, exist_ok=True)
    tmp = DATA_FILE.with_suffix(".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(DATA_FILE)


def merge(data, ev, now):
    events = data["events"]
    ev["status"] = "Confirmed"
    old = events.get(ev["key"])
    if not old:
        ev.update(version=1, first_seen=now, last_updated=now, last_seen=now)
        events[ev["key"]] = ev
        print(f"  NEW      {ev['start']}  {ev['name'][:70]}")
        return
    changed = any(old.get(f) != ev.get(f) for f in TRACKED)
    ev["first_seen"] = old.get("first_seen", now)
    ev["last_seen"] = now
    ev["version"] = old.get("version", 1) + (1 if changed else 0)
    ev["last_updated"] = now if changed else old.get("last_updated", now)
    events[ev["key"]] = {**old, **ev}
    if changed:
        print(f"  CHANGED  {ev['start']}  {ev['name'][:70]}")


# ---------------------------------------------------------------- outputs
def body_html(e):
    def row(label, value):
        return f"<b>{label}:</b> {value}<br>" if value else ""

    def link(label, url):
        return f'<b>{label}:</b> <a href="{url}">{url}</a><br>' if url else ""

    note = ""
    if e["status"] != "Confirmed":
        note = ('<p style="color:#b00"><b>Note:</b> this event is no longer listed on the '
                'WTC Binh Duong website. Check if it was moved or cancelled.</p>')
    dates = e["start"] if e["start"] == e["end"] else f"{e['start']} to {e['end']}"
    desc = "".join(f"<p>{p}</p>" for p in e.get("description", "").replace("<", "&lt;").split("\n\n") if p)
    desc = desc.replace("\n", "<br>")
    return (
        note
        + row("Dates", dates) + row("Hours", e.get("hours")) + row("Venue", e.get("venue"))
        + row("Hall", e.get("hall")) + row("Address", e.get("address"))
        + row("Vietnamese name", e.get("name_vn") if e.get("name_vn") != e["name"] else "")
        + link("Event page", e.get("url")) + link("Website", e.get("website"))
        + link("Registration", e.get("registration")) + link("Brochure", e.get("brochure"))
        + link("Map", e.get("map")) + row("Status", e["status"])
        + "<hr>" + desc
    )


def body_text(e):
    lines = []
    if e["status"] != "Confirmed":
        lines.append("NOTE: no longer listed on the WTC Binh Duong website. "
                     "Check if it was moved or cancelled.\n")
    dates = e["start"] if e["start"] == e["end"] else f"{e['start']} to {e['end']}"
    for label, key in (("Hours", "hours"), ("Venue", "venue"), ("Hall", "hall"),
                       ("Address", "address"), ("Event page", "url"), ("Website", "website"),
                       ("Registration", "registration"), ("Brochure", "brochure"),
                       ("Map", "map"), ("Status", "status")):
        if e.get(key):
            lines.append(f"{label}: {e[key]}")
    lines.insert(0 if e["status"] == "Confirmed" else 1, f"Dates: {dates}")
    if e.get("name_vn") and e["name_vn"] != e["name"]:
        lines.append(f"Vietnamese name: {e['name_vn']}")
    if e.get("description"):
        lines.append("\n" + e["description"])
    return "\n".join(lines)


def location(e):
    venue = e["venue"] + (f" - {e['hall']}" if e.get("hall") else "")
    return f"{venue}, {e['address']}"


def ics_escape(s):
    return (s or "").replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


def ics_fold(line):
    out, b = [], line.encode("utf-8")
    while len(b) > 74:
        cut = 74
        while (b[cut] & 0xC0) == 0x80:  # don't split a UTF-8 character
            cut -= 1
        out.append(b[:cut].decode("utf-8"))
        b = b" " + b[cut:]
    out.append(b.decode("utf-8"))
    return "\r\n".join(out)


def write_ics(events):
    lines = [
        "BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//JTM Asia//WTC Binh Duong//EN",
        "CALSCALE:GREGORIAN", "METHOD:PUBLISH",
        "X-WR-CALNAME:WTC Binh Duong Trade Fairs", "X-WR-TIMEZONE:Asia/Ho_Chi_Minh",
        "REFRESH-INTERVAL;VALUE=DURATION:P1D", "X-PUBLISHED-TTL:P1D",
    ]
    for e in events:
        start = date.fromisoformat(e["start"])
        end = date.fromisoformat(e["end"]) + timedelta(days=1)
        stamp = e["last_updated"].replace("-", "").replace(":", "")[:15] + "Z"
        lines += [
            "BEGIN:VEVENT",
            f"UID:wtcbd-{e['key']}@jtmasia.com",
            f"DTSTAMP:{stamp}",
            f"SEQUENCE:{e['version']}",
            f"DTSTART;VALUE=DATE:{start:%Y%m%d}",
            f"DTEND;VALUE=DATE:{end:%Y%m%d}",
            f"SUMMARY:{ics_escape(e['name'])}",
            f"LOCATION:{ics_escape(location(e))}",
            f"DESCRIPTION:{ics_escape(body_text(e))}",
            f"URL:{e['url']}",
            "TRANSP:TRANSPARENT",
            "STATUS:" + ("CONFIRMED" if e["status"] == "Confirmed" else "TENTATIVE"),
            "END:VEVENT",
        ]
    lines.append("END:VCALENDAR")
    ICS_FILE.write_text("\r\n".join(ics_fold(l) for l in lines) + "\r\n", encoding="utf-8")


def write_pa_json(events, now):
    """Upcoming events only, in the same shape as the SECC events.json."""
    out = []
    for e in events:
        if date.fromisoformat(e["end"]) < TODAY:
            continue
        end_excl = date.fromisoformat(e["end"]) + timedelta(days=1)
        out.append({
            "key": e["key"],
            "name": e["name"],
            "start": e["start"],
            "end": e["end"],
            "start_time": f"{e['start']}T00:00:00",
            "end_time": f"{end_excl.isoformat()}T00:00:00",
            "hours": e.get("hours", ""),
            "venue": e["venue"],
            "hall": e.get("hall", ""),
            "location": location(e),
            "url": e["url"],
            "status": e["status"],
            "version": e["version"],
            "body_html": body_html(e),
        })
    PA_FILE.write_text(json.dumps({"generated": now, "source": LIST_URL, "events": out},
                                  ensure_ascii=False, indent=2), encoding="utf-8")
    return len(out)


def write_xlsx(events):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    ws = wb.active
    ws.title = "WTC Binh Duong Trade Fairs"
    cols = [("Name", "name", 50), ("Start", "start", 12), ("End", "end", 12),
            ("Hours", "hours", 16), ("Venue", "venue", 40), ("Hall", "hall", 9),
            ("Status", "status", 22), ("Vietnamese name", "name_vn", 50),
            ("Event page", "url", 45), ("Website", "website", 35),
            ("Registration", "registration", 35), ("Brochure", "brochure", 35),
            ("Description", "description", 80), ("Last updated", "last_updated", 20)]
    ws.append([c[0] for c in cols])
    for e in events:
        ws.append([e.get(k, "") for _, k, _ in cols])
    for i, (_, _, w) in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.cell(1, i).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, i).fill = PatternFill("solid", fgColor="1F4E78")
    ws.freeze_panes = "B2"
    if events:
        ref = f"A1:{ws.cell(1, len(cols)).column_letter}{len(events) + 1}"
        t = Table(displayName="WTCEvents", ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
    wb.save(XLSX_FILE)


# ---------------------------------------------------------------- main
def main():
    now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    data = load_data()
    events = data["events"]

    cutoff = HISTORY_FROM
    if events:  # later runs: only look back 60 days, older events don't change
        cutoff = max(HISTORY_FROM, TODAY - timedelta(days=60))
    print(f"Reading listing back to {cutoff} ...")
    items = scan_listing(cutoff)
    if not items:
        print("ERROR: 0 events found. The WTC website layout may have changed "
              "or the request was blocked. Nothing was overwritten.")
        sys.exit(1)
    print(f"{len(items)} events in range. Reading event pages ...")

    seen, failed = set(), 0
    for i, item in enumerate(items, 1):
        key = slug_of(item["url"])
        seen.add(key)
        old = events.get(key)
        # Past events already saved are not re-read.
        if old and date.fromisoformat(old["end"]) < TODAY - timedelta(days=7):
            continue
        print(f"[{i}/{len(items)}] {item['name_vn'][:70]}")
        try:
            ev = fetch_event(item)
        except Exception as e:  # one bad page must not stop the run
            failed += 1
            print(f"  SKIPPED (error): {e}")
            continue
        if ev:
            merge(data, ev, now)
            save_data(data)  # crash-safe: saved after every event

    # Upcoming events that vanished from the site: keep, mark as tentative.
    for key, e in events.items():
        if key in seen or date.fromisoformat(e["end"]) < TODAY:
            continue
        if date.fromisoformat(e["start"]) < cutoff:  # older than the part of the list we read
            continue
        if e.get("status") == "Confirmed":
            e["status"] = f"Not listed since {TODAY.isoformat()}"
            e["version"] = e.get("version", 1) + 1
            e["last_updated"] = now
            print(f"  NOT LISTED  {e['start']}  {e['name'][:70]}")

    data["last_run"] = now
    save_data(data)

    ordered = sorted(events.values(), key=lambda e: (e["start"], e["name"]))
    DOCS.mkdir(exist_ok=True)
    write_ics(ordered)
    upcoming = write_pa_json(ordered, now)
    write_xlsx(ordered)
    print(f"Done. {len(ordered)} events in calendar, {upcoming} upcoming. "
          f"{failed} pages failed.")
    if failed and failed == len(items):
        sys.exit(1)


if __name__ == "__main__":
    main()
